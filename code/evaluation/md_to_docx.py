"""
Konversi Kerangka_AN_Klasifikasi_GSS_ML.md → Kerangka_AN_Klasifikasi_GSS_ML.docx
"""
from __future__ import annotations
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT     = r"D:\1. Important\Work\Bank Indonesia\DSta-DSMF\Green Bond Classification"
MD_PATH  = os.path.join(ROOT, "Kerangka_AN_Klasifikasi_GSS_ML.md")
OUT_PATH = os.path.join(ROOT, "Kerangka_AN_Klasifikasi_GSS_ML.docx")


# ---------------------------------------------------------------------------
# Inline-markdown tokeniser: bold / italic / code / link → runs
# ---------------------------------------------------------------------------
_INLINE = re.compile(
    r'(\*\*(?:[^*]|\*(?!\*))+\*\*'   # **bold**
    r'|\*(?!\*)(?:[^*])+(?<!\*)\*'   # *italic*  (not **)
    r'|`[^`]+`'                        # `code`
    r'|\[[^\]]+\]\([^)]*\))'          # [text](url)
)

def _add_inline(para, text: str):
    pos = 0
    for m in _INLINE.finditer(text):
        # plain text before match
        if m.start() > pos:
            para.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            run = para.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("`"):
            run = para.add_run(tok[1:-1])
            run.font.name = "Courier New"; run.font.size = Pt(9)
        elif tok.startswith("["):
            link_text = re.match(r'\[([^\]]+)\]', tok).group(1)
            run = para.add_run(link_text)
            run.font.color.rgb = RGBColor(0x00, 0x56, 0xB2)
        else:  # *italic*
            run = para.add_run(tok[1:-1]); run.italic = True
        pos = m.end()
    if pos < len(text):
        para.add_run(text[pos:])


# ---------------------------------------------------------------------------
# Table helpers
# ---------------------------------------------------------------------------
_SEP_ROW = re.compile(r'^\|[-:| ]+\|$')

def _parse_table_rows(raw_rows: list[str]) -> list[list[str]]:
    parsed = []
    for row in raw_rows:
        if _SEP_ROW.match(row.strip()):
            continue
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    return parsed

def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def _flush_table(doc: Document, raw_rows: list[str]):
    parsed = _parse_table_rows(raw_rows)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    tbl   = doc.add_table(rows=len(parsed), cols=ncols)
    tbl.style = "Table Grid"
    for ri, row in enumerate(parsed):
        for ci in range(ncols):
            cell_md = row[ci] if ci < len(row) else ""
            cell    = tbl.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            # strip **bold** markers, remember if cell was bold
            bold_cell = bool(re.search(r'\*\*', cell_md))
            clean = cell_md.replace("**", "")
            run = para.add_run(clean)
            if ri == 0 or bold_cell:
                run.bold = True
            # align numeric-ish right, header and first col left
            if ri > 0 and ci > 0:
                try:
                    float(clean.replace(",", ".").replace("%", "").replace("Rp", "").strip())
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                except ValueError:
                    pass
        # header row shading
        if ri == 0:
            for ci in range(ncols):
                _set_cell_bg(tbl.cell(ri, ci), "D6E4F0")
    doc.add_paragraph()   # spacing after table


# ---------------------------------------------------------------------------
# Main converter
# ---------------------------------------------------------------------------
def convert(md_path: str, out_path: str):
    doc = Document()

    # Page margins
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3.0)
        sec.right_margin  = Cm(2.5)

    # Default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    with open(md_path, encoding="utf-8") as f:
        lines = [l.rstrip("\n") for l in f]

    in_frontmatter = False
    in_table       = False
    table_rows: list[str] = []

    for raw in lines:
        stripped = raw.strip()

        # ── YAML front matter ────────────────────────────────────────────
        if not in_frontmatter and stripped == "---" and lines.index(raw) == 0:
            in_frontmatter = True; continue
        if in_frontmatter:
            if stripped == "---": in_frontmatter = False
            continue

        # ── Table rows ──────────────────────────────────────────────────
        if stripped.startswith("|"):
            if not in_table:
                in_table = True; table_rows = []
            table_rows.append(stripped)
            continue
        else:
            if in_table:
                _flush_table(doc, table_rows)
                in_table = False; table_rows = []

        # ── Horizontal rule ─────────────────────────────────────────────
        if re.match(r'^-{3,}$', stripped) or stripped in ("***", "___"):
            doc.add_paragraph()
            continue

        # ── Empty line ──────────────────────────────────────────────────
        if not stripped:
            continue

        # ── Headings ────────────────────────────────────────────────────
        m = re.match(r'^(#{1,3})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            para  = doc.add_heading(level=level)
            para.clear()               # remove run added by add_heading
            _add_inline(para, m.group(2))
            continue

        # ── Blockquote ──────────────────────────────────────────────────
        if stripped.startswith("> "):
            para = doc.add_paragraph(style="Quote")
            _add_inline(para, stripped[2:])
            continue

        # ── Numbered list ───────────────────────────────────────────────
        m = re.match(r'^\d+\.\s+(.*)', stripped)
        if m:
            para = doc.add_paragraph(style="List Number")
            _add_inline(para, m.group(1))
            continue

        # ── Bullet list ─────────────────────────────────────────────────
        if stripped.startswith("- ") or stripped.startswith("* "):
            para = doc.add_paragraph(style="List Bullet")
            _add_inline(para, stripped[2:])
            continue

        # ── Normal paragraph ────────────────────────────────────────────
        para = doc.add_paragraph()
        _add_inline(para, stripped)

    # flush trailing table
    if in_table:
        _flush_table(doc, table_rows)

    doc.save(out_path)
    print(f"Tersimpan: {out_path}")


if __name__ == "__main__":
    convert(MD_PATH, OUT_PATH)
