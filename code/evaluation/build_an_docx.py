"""
Pembangun dokumen Analytical Note (AN) yang PROPER & siap paparan:
  - Halaman sampul (cover) ber-branding BI / DSta-DSMF
  - Daftar isi otomatis (field TOC — tekan F9 di Word bila perlu refresh)
  - Grafik PNG TER-EMBED inline di seksi yang relevan + caption "Grafik N."
  - Tabel data ter-format (header berwarna, angka rata-kanan)
  - Heading berwarna biru BI, footer nomor halaman, ukuran kertas A4

Sumber teks  : Deliverable AN GSS/Kerangka_AN_Klasifikasi_GSS_ML.md
Sumber grafik: Deliverable AN GSS/03_Grafik/*.png
Output       : Deliverable AN GSS/Kerangka_AN_Klasifikasi_GSS_ML.docx
               Deliverable AN GSS/01_Dokumen/Analytical_Note_Klasifikasi_GSS_EBUS_Korporasi.docx

Jalankan dari root repo (pastikan grafik sudah diregenerasi via make_deliverable.py):
  python code/evaluation/build_an_docx.py
"""
from __future__ import annotations
import os
import re
import shutil

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT      = r"D:\1. Important\Work\Bank Indonesia\DSta-DSMF\Green Bond Classification"
OUT_DIR   = os.path.join(ROOT, "Deliverable AN GSS")
MD_PATH   = os.path.join(OUT_DIR, "Kerangka_AN_Klasifikasi_GSS_ML.md")
CHART_DIR = os.path.join(OUT_DIR, "03_Grafik")
DOCX_SRC  = os.path.join(OUT_DIR, "Kerangka_AN_Klasifikasi_GSS_ML.docx")
DOCX_DLV  = os.path.join(OUT_DIR, "01_Dokumen",
                         "Analytical_Note_Klasifikasi_GSS_EBUS_Korporasi.docx")

# Warna
BI_BLUE   = RGBColor(0x00, 0x56, 0xB2)
BI_DARK   = RGBColor(0x00, 0x3A, 0x78)
GREY_TXT  = RGBColor(0x55, 0x55, 0x55)
HDR_FILL  = "D6E4F0"

CONTENT_WIDTH_IN = 6.0   # lebar konten A4 (margin 3cm/2,5cm)

# ---------------------------------------------------------------------------
# Penempatan grafik: setelah baris yang memuat anchor → sisipkan chart + caption
# ---------------------------------------------------------------------------
CHART_ANCHORS: list[tuple[str, list[tuple[str, str]]]] = [
    ("dari semua yang dilabeli GSS oleh sistem", [
        ("03_evaluasi_rule_vs_ml.png",
         "Evaluasi klasifikasi GSS — rule-based vs ML semantik (gold set 93 dokumen)."),
    ]),
    ("Social Bond mendominasi baik dari segi jumlah", [
        ("01_sensus_pasar_per_kelas.png",
         "Sensus pasar GSS korporasi per kelas — jumlah instrumen & outstanding."),
        ("02_gap_labeled_vs_universe.png",
         "Gap pelabelan: 82 GSS berlabel dari 1.437 instrumen EBUS korporasi."),
    ]),
    ("hadir di 6 dari 9 instrumen yang terpetakan", [
        ("04_dekomposisi_sektoral.png",
         "Dekomposisi sektoral use-of-proceeds GSS korporasi (gold set N=27)."),
    ]),
    ("Tidak satu pun dari 198 dokumen yang berhasil dibaca", [
        ("05_cakupan_pemindaian_semesta.png",
         "Cakupan pemindaian semesta EBUS & estimasi GSS tidak berlabel."),
    ]),
    ("06_statistik_deskriptif_gss.csv`, dihasilkan", [
        ("06_tren_penerbitan_gss.png",
         "Tren penerbitan EBUS GSS korporasi per tahun (2022–2026)."),
        ("07_profil_jatuh_tempo.png",
         "Profil jatuh tempo outstanding GSS korporasi (maturity ladder)."),
        ("08_konsentrasi_dan_rating.png",
         "Konsentrasi penerbit (HHI) & sebaran rating instrumen GSS."),
    ]),
]


# ---------------------------------------------------------------------------
# Inline markdown → runs
# ---------------------------------------------------------------------------
_INLINE = re.compile(
    r'(\*\*(?:[^*]|\*(?!\*))+\*\*'
    r'|\*(?!\*)(?:[^*])+(?<!\*)\*'
    r'|`[^`]+`'
    r'|\[[^\]]+\]\([^)]*\))'
)

def _add_inline(para, text: str):
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            para.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            run = para.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("`"):
            run = para.add_run(tok[1:-1])
            run.font.name = "Consolas"; run.font.size = Pt(9.5)
            run.font.color.rgb = BI_DARK
        elif tok.startswith("["):
            link_text = re.match(r'\[([^\]]+)\]', tok).group(1)
            run = para.add_run(link_text)
            run.font.color.rgb = BI_BLUE; run.underline = True
        else:
            run = para.add_run(tok[1:-1]); run.italic = True
        pos = m.end()
    if pos < len(text):
        para.add_run(text[pos:])


# ---------------------------------------------------------------------------
# Tabel
# ---------------------------------------------------------------------------
_SEP_ROW = re.compile(r'^\|[-:| ]+\|$')

def _set_cell_bg(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def _flush_table(doc, raw_rows):
    parsed = []
    for row in raw_rows:
        if _SEP_ROW.match(row.strip()):
            continue
        parsed.append([c.strip() for c in row.strip().strip("|").split("|")])
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    tbl = doc.add_table(rows=len(parsed), cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(parsed):
        for ci in range(ncols):
            cell_md = row[ci] if ci < len(row) else ""
            cell = tbl.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            bold_cell = bool(re.search(r'\*\*', cell_md))
            clean = cell_md.replace("**", "")
            run = para.add_run(clean)
            run.font.size = Pt(9.5)
            if ri == 0:
                run.bold = True; run.font.color.rgb = BI_DARK
            elif bold_cell:
                run.bold = True
            if ri > 0 and ci > 0:
                try:
                    float(clean.replace(",", ".").replace("%", "")
                          .replace("Rp", "").strip())
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                except ValueError:
                    pass
        if ri == 0:
            for ci in range(ncols):
                _set_cell_bg(tbl.cell(ri, ci), HDR_FILL)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Grafik
# ---------------------------------------------------------------------------
def _insert_chart(doc, fig_no: int, filename: str, caption: str) -> int:
    path = os.path.join(CHART_DIR, filename)
    if not os.path.exists(path):
        print(f"  ! grafik tak ada, dilewati: {filename}")
        return fig_no
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(CONTENT_WIDTH_IN))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = cap.add_run(f"Grafik {fig_no}. ")
    r1.bold = True; r1.font.size = Pt(9); r1.font.color.rgb = BI_BLUE
    r2 = cap.add_run(caption)
    r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = GREY_TXT
    doc.add_paragraph()
    print(f"  + Grafik {fig_no}: {filename}")
    return fig_no + 1


# ---------------------------------------------------------------------------
# Cover + TOC + footer
# ---------------------------------------------------------------------------
def _hrule(para):
    pPr = para._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "0056B2")
    pbdr.append(bottom); pPr.append(pbdr)

def _build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BANK INDONESIA")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = BI_BLUE
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Departemen Statistik (DSta) · DSMF")
    r.font.size = Pt(11); r.font.color.rgb = GREY_TXT

    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); _hrule(p)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ANALYTICAL NOTE")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREY_TXT

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Klasifikasi Efek Bersifat Utang & Sukuk (EBUS) Korporasi\n"
                  "ke dalam Kategori Green, Social, Sustainability (GSS)")
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = BI_DARK

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Pendekatan Machine Learning Tergrounding Taksonomi "
                  "POJK 18/2023, ICMA, dan Kerangka DJPPR")
    r.italic = True; r.font.size = Pt(11.5); r.font.color.rgb = GREY_TXT

    p = doc.add_paragraph(); _hrule(p)

    for _ in range(8):
        doc.add_paragraph()

    for label, value in [
        ("Unit", "Bank Indonesia — DSta / DSMF"),
        ("Periode data", "Listing IDX per 18 Juni 2026"),
        ("Status", "Draft — siap paparan pemangku kepentingan"),
        ("Sifat", "Internal / terbatas"),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}: "); r1.bold = True; r1.font.size = Pt(10.5)
        r2 = p.add_run(value); r2.font.size = Pt(10.5); r2.font.color.rgb = GREY_TXT

    doc.add_page_break()

def _add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run("Daftar Isi")
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BI_BLUE
    doc.add_paragraph()

    par = doc.add_paragraph()
    run = par.add_run()
    fldBegin = OxmlElement("w:fldChar"); fldBegin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fldSep = OxmlElement("w:fldChar"); fldSep.set(qn("w:fldCharType"), "separate")
    fldText = OxmlElement("w:t")
    fldText.text = "Klik kanan → Update Field (F9) untuk memuat daftar isi."
    fldEnd = OxmlElement("w:fldChar"); fldEnd.set(qn("w:fldCharType"), "end")
    for el in (fldBegin, instr, fldSep, fldText, fldEnd):
        run._r.append(el)
    doc.add_page_break()

def _add_page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fldBegin = OxmlElement("w:fldChar"); fldBegin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fldEnd = OxmlElement("w:fldChar"); fldEnd.set(qn("w:fldCharType"), "end")
    for el in (fldBegin, instr, fldEnd):
        run._r.append(el)
    run.font.size = Pt(9); run.font.color.rgb = GREY_TXT


# ---------------------------------------------------------------------------
# Styling heading
# ---------------------------------------------------------------------------
def _style_headings(doc):
    for name, size, color in [
        ("Heading 1", 16, BI_BLUE),
        ("Heading 2", 13, BI_DARK),
        ("Heading 3", 11.5, BI_DARK),
    ]:
        st = doc.styles[name]
        st.font.size = Pt(size); st.font.color.rgb = color
        st.font.bold = True; st.font.name = "Calibri"


# ---------------------------------------------------------------------------
# Konverter utama
# ---------------------------------------------------------------------------
def build():
    doc = Document()

    # Kertas A4 + margin
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"; style.font.size = Pt(11)
    _style_headings(doc)

    _build_cover(doc)
    _add_toc(doc)
    _add_page_number_footer(doc)

    with open(MD_PATH, encoding="utf-8") as f:
        lines = [l.rstrip("\n") for l in f]

    # Mulai body dari "## 0." (lewati metadata judul atas)
    start = 0
    for i, l in enumerate(lines):
        if re.match(r'^##\s*0\.', l.strip()):
            start = i; break
    lines = lines[start:]

    in_table = False
    table_rows: list[str] = []
    fig_no = 1

    def _maybe_charts(text_line: str):
        nonlocal fig_no
        for anchor, charts in CHART_ANCHORS:
            if anchor in text_line:
                for fname, cap in charts:
                    fig_no = _insert_chart(doc, fig_no, fname, cap)

    for raw in lines:
        stripped = raw.strip()

        if stripped.startswith("|"):
            if not in_table:
                in_table = True; table_rows = []
            table_rows.append(stripped)
            continue
        else:
            if in_table:
                _flush_table(doc, table_rows)
                in_table = False; table_rows = []

        if re.match(r'^-{3,}$', stripped) or stripped in ("***", "___"):
            continue
        if not stripped:
            continue

        m = re.match(r'^(#{1,3})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            para = doc.add_heading(level=level)
            para.clear()
            _add_inline(para, m.group(2))
            continue

        if stripped.startswith("> "):
            para = doc.add_paragraph(style="Quote")
            _add_inline(para, stripped[2:])
            continue

        m = re.match(r'^\d+\.\s+(.*)', stripped)
        if m:
            para = doc.add_paragraph(style="List Number")
            _add_inline(para, m.group(1))
            _maybe_charts(stripped)
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            para = doc.add_paragraph(style="List Bullet")
            _add_inline(para, stripped[2:])
            continue

        para = doc.add_paragraph()
        _add_inline(para, stripped)
        # sisipkan grafik setelah baris caption "*Sumber: ...*"
        _maybe_charts(stripped)

    if in_table:
        _flush_table(doc, table_rows)

    doc.save(DOCX_SRC)
    os.makedirs(os.path.dirname(DOCX_DLV), exist_ok=True)
    shutil.copy2(DOCX_SRC, DOCX_DLV)
    print(f"\nTersimpan ({fig_no-1} grafik ter-embed):")
    print(f"  {os.path.relpath(DOCX_SRC, ROOT)}")
    print(f"  {os.path.relpath(DOCX_DLV, ROOT)}")


if __name__ == "__main__":
    build()
